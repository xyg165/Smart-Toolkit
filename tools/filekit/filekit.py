#!/usr/bin/env python3
"""
filekit —— 文件整理工具箱（rename + classify 合体版）
================================================================
三种模式：
  rename     只按文件名规范（清理 + 类型前缀 + 时间戳）—— 零依赖
  classify   读文档内容自动归类（38 类）—— 需依赖
  auto       智能：能读到内容就归类，读不到就只按文件名
  rollback   按日志回滚

用法：
  python filekit.py rename    --dir ./素材 --pad --apply
  python filekit.py classify  --dir ./文档 --verbose
  python filekit.py auto      --dir ./混合文件夹 --apply
  python filekit.py rollback  ./_filekit_log_20261007_160000.json

命名规范：
  [类型]_[分类]_[原名]_[版本]_[YYMMDDHHMM].ext
   DOC    合同   采购合同   v1    2609151030
"""
import argparse, os, re, json, sys, datetime, subprocess, shutil
from pathlib import Path

VERSION = "3.3"
MAX_CHARS = 12000          # classify 模式每个文件最多读多少字

# ============================================================================
# 一、类型前缀
# ============================================================================
TYPE_PREFIX = {
    '.jpg': 'IMG', '.jpeg': 'IMG', '.png': 'IMG', '.gif': 'IMG', '.webp': 'IMG',
    '.bmp': 'IMG', '.svg': 'IMG', '.ico': 'IMG', '.heic': 'IMG', '.tif': 'IMG', '.tiff': 'IMG',
    '.doc': 'DOC', '.docx': 'DOC', '.md': 'DOC', '.txt': 'DOC', '.rtf': 'DOC', '.odt': 'DOC',
    '.xls': 'XLS', '.xlsx': 'XLS', '.csv': 'XLS', '.ods': 'XLS',
    '.ppt': 'PPT', '.pptx': 'PPT', '.key': 'PPT', '.odp': 'PPT',
    '.pdf': 'PDF',
    '.mp3': 'AUD', '.wav': 'AUD', '.m4a': 'AUD', '.flac': 'AUD', '.aac': 'AAC',
    '.mp4': 'VID', '.mov': 'VID', '.avi': 'VID', '.mkv': 'VID', '.webm': 'VID',
    '.zip': 'ZIP', '.rar': 'ZIP', '.7z': 'ZIP', '.tar': 'ZIP', '.gz': 'ZIP',
    '.psd': 'IMG', '.ai': 'IMG', '.sketch': 'IMG',
}
PREFIX_RE = re.compile(r'^(IMG|DOC|XLS|PPT|PDF|AUD|AAC|VID|ZIP|FILE)[_\-]', re.I)

# ============================================================================
# 二、内容分类库（38 类 / 479 关键词）
#    strong = 强特征词（3 分）  weak = 一般词（1 分）
# ============================================================================
CATEGORIES = {
    # ---- 商务经营 ----
    "合同": {
        "strong": ["甲方", "乙方", "丙方", "本合同", "违约责任", "合同编号", "签订地点",
                   "合同期限", "合同价款", "合同生效", "解除本合同"],
        "weak": ["合同", "协议", "签署", "盖章", "生效", "标的", "条款", "争议解决", "履约"],
    },
    "保密协议": {
        "strong": ["保密义务", "保密信息", "商业秘密", "不得披露", "保密期限", "泄密"],
        "weak": ["保密", "机密", "NDA", "保密条款"],
    },
    "报价单": {
        "strong": ["报价单", "报价明细", "报价有效期", "单价", "合计金额", "优惠价"],
        "weak": ["报价", "总价", "折扣", "含税", "不含税", "供货周期"],
    },
    "招投标": {
        "strong": ["招标文件", "投标文件", "投标人", "招标人", "评标", "中标", "开标",
                   "投标保证金", "招标公告"],
        "weak": ["招标", "投标", "标书", "标段", "资质要求"],
    },
    "采购订单": {
        "strong": ["采购订单", "采购合同", "供应商名称", "到货日期", "请购单", "采购申请"],
        "weak": ["采购", "供应商", "订货", "交货期", "数量", "PO"],
    },
    "发票": {
        "strong": ["发票代码", "发票号码", "纳税人识别号", "价税合计", "开票日期",
                   "销售方", "购买方", "增值税", "税率"],
        "weak": ["发票", "税额", "收款人", "复核", "合计"],
    },
    "付款凭证": {
        "strong": ["付款凭证", "银行回单", "转账凭证", "收款凭证", "电子回单", "交易流水"],
        "weak": ["付款", "收款", "汇款", "账号", "开户行", "金额"],
    },
    "客户资料": {
        "strong": ["客户名称", "客户需求", "客户联系人", "客户档案", "跟进记录"],
        "weak": ["客户", "联系人", "商务", "合作意向", "拜访"],
    },
    "销售合同": {
        "strong": ["销售合同", "供货方", "采购方", "订货方", "产品清单", "发货"],
        "weak": ["销售", "订单", "货款", "退货", "售后"],
    },
    # ---- 项目产品 ----
    "需求文档": {
        "strong": ["需求文档", "产品需求", "功能需求", "用户故事", "验收标准", "需求描述",
                   "PRD", "需求编号", "优先级"],
        "weak": ["需求", "功能", "用例", "功能清单", "用户场景"],
    },
    "方案": {
        "strong": ["实施方案", "技术方案", "总体方案", "总体思路", "建设方案", "推进计划",
                   "总体目标", "阶段目标"],
        "weak": ["方案", "规划", "计划书", "策略", "落地", "举措"],
    },
    "项目计划": {
        "strong": ["项目计划", "WBS", "里程碑", "任务分解", "资源分配", "进度计划",
                   "关键路径", "工期"],
        "weak": ["计划", "进度", "任务", "排期", "甘特"],
    },
    "测试报告": {
        "strong": ["测试报告", "测试用例", "测试结果", "缺陷", "通过率", "回归测试",
                   "测试环境", "测试结论"],
        "weak": ["测试", "验证", "BUG", "用例", "性能", "压力"],
    },
    "验收文档": {
        "strong": ["验收报告", "验收标准", "交付验收", "验收结论", "移交清单", "验收单"],
        "weak": ["验收", "交付", "签收", "确认", "交接"],
    },
    "项目周报": {
        "strong": ["项目周报", "本周工作", "下周计划", "工作进展", "完成情况", "周报"],
        "weak": ["周报", "月报", "进展", "汇报", "待办"],
    },
    "产品文档": {
        "strong": ["产品文档", "版本说明", "发布说明", "产品路线", "迭代计划", "功能清单"],
        "weak": ["产品", "版本", "发布", "迭代", "更新"],
    },
    # ---- 技术研发 ----
    "技术文档": {
        "strong": ["接口文档", "API文档", "系统架构", "数据字典", "部署文档", "接口说明",
                   "请求参数", "响应参数", "技术架构"],
        "weak": ["接口", "架构", "部署", "模块", "服务", "配置"],
    },
    "设计文档": {
        "strong": ["概要设计", "详细设计", "设计说明", "数据库设计", "UML", "时序图",
                   "类图", "状态机"],
        "weak": ["设计", "方案设计", "逻辑", "流程", "结构"],
    },
    "说明书": {
        "strong": ["使用说明书", "操作手册", "安装说明", "技术参数", "规格参数",
                   "维护保养", "注意事项", "产品规格"],
        "weak": ["说明书", "手册", "操作步骤", "安装", "使用", "参数"],
    },
    "图纸": {
        "strong": ["装配图", "零件图", "工程图", "尺寸公差", "图纸编号", "CAD", "三视图"],
        "weak": ["图纸", "图样", "尺寸", "标注", "比例"],
    },
    "知识产权": {
        "strong": ["专利", "专利申请", "著作权", "商标注册", "知识产权", "发明人",
                   "申请号", "权利要求书"],
        "weak": ["发明", "实用新型", "软著", "授权", "登记"],
    },
    "技术标准": {
        "strong": ["国家标准", "行业标准", "企业标准", "技术规范", "GB/T", "技术条件",
                   "检验标准"],
        "weak": ["标准", "规范", "规程", "要求", "指标"],
    },
    # ---- 行政人事 ----
    "规章制度": {
        "strong": ["管理制度", "管理办法", "实施细则", "本办法", "本制度", "奖惩",
                   "考核办法", "职责分工"],
        "weak": ["制度", "规定", "流程", "职责", "考核", "执行"],
    },
    "通知公告": {
        "strong": ["特此通知", "各部门", "通知如下", "公告", "现将有关事项通知"],
        "weak": ["通知", "告知", "安排", "要求", "传达"],
    },
    "会议纪要": {
        "strong": ["会议纪要", "参会人员", "会议时间", "会议地点", "主持人", "记录人",
                   "会议决议", "行动项", "议题"],
        "weak": ["会议", "纪要", "决议", "讨论", "与会"],
    },
    "简历": {
        "strong": ["个人简历", "求职意向", "教育背景", "工作经历", "项目经验",
                   "毕业院校", "应聘岗位"],
        "weak": ["简历", "求职", "应聘", "技能", "特长"],
    },
    "人事档案": {
        "strong": ["劳动合同", "入职登记", "离职申请", "转正申请", "绩效考核",
                   "员工档案", "社保", "薪资"],
        "weak": ["入职", "离职", "转正", "考勤", "薪酬", "员工"],
    },
    "培训材料": {
        "strong": ["培训材料", "培训课程", "讲义", "课后练习", "培训计划", "考核试题"],
        "weak": ["培训", "课程", "教学", "学习", "知识点", "考核"],
    },
    "工作总结": {
        "strong": ["工作总结", "年度总结", "半年总结", "个人总结", "述职报告"],
        "weak": ["总结", "回顾", "成绩", "不足", "展望"],
    },
    # ---- 财务 ----
    "财务报表": {
        "strong": ["资产负债表", "利润表", "现金流量表", "所有者权益", "会计科目",
                   "借方", "贷方", "期末余额"],
        "weak": ["财务", "报表", "余额", "核算", "账务"],
    },
    "预算": {
        "strong": ["预算编制", "费用预算", "预算执行", "预算方案", "年度预算", "预算科目"],
        "weak": ["预算", "费用", "支出", "成本", "额度"],
    },
    "报销单": {
        "strong": ["费用报销", "差旅报销", "报销单", "报销凭证", "交通费", "住宿费"],
        "weak": ["报销", "差旅", "票据", "发票粘贴", "审批"],
    },
    # ---- 内容运营 ----
    "内容脚本": {
        "strong": ["口播", "分镜", "脚本", "开场白", "结尾", "选题方向", "话术",
                   "内容大纲", "视频脚本"],
        "weak": ["脚本", "文案", "选题", "标题", "内容"],
    },
    "运营数据": {
        "strong": ["播放量", "完播率", "粉丝增长", "转化率", "阅读量", "数据复盘",
                   "互动率", "涨粉"],
        "weak": ["数据", "统计", "分析", "报表", "指标"],
    },
    "素材清单": {
        "strong": ["素材清单", "封面图", "配图", "音乐素材", "字体素材", "版权说明"],
        "weak": ["素材", "图片", "视频素材", "音频", "资源"],
    },
    # ---- 法律与其他 ----
    "法律文书": {
        "strong": ["起诉状", "答辩状", "律师函", "判决书", "仲裁申请", "诉讼请求",
                   "委托代理", "法律意见书"],
        "weak": ["法律", "诉讼", "仲裁", "律师", "法院"],
    },
    "资质证书": {
        "strong": ["营业执照", "资质证书", "许可证", "认证证书", "证书编号",
                   "统一社会信用代码", "资质等级"],
        "weak": ["资质", "证书", "认证", "许可", "年检"],
    },
    "分析报告": {
        "strong": ["分析报告", "调研报告", "可行性研究", "市场分析", "研究结论",
                   "数据来源", "结论与建议"],
        "weak": ["分析", "调研", "研究", "结论", "建议", "评估"],
    },
}
_WEIGHTED = {c: [(k, 3.0) for k in v.get("strong", [])] + [(k, 1.0) for k in v.get("weak", [])]
             for c, v in CATEGORIES.items()}

# ============================================================================
# 环境自检 —— 告诉用户缺什么、怎么装（而不是静默失败）
# ============================================================================
DEP_HINTS = {
    'rapidocr':       ('rapidocr-onnxruntime（图片 OCR，纯 Python 免安装）', 'pip install rapidocr-onnxruntime'),
    'tesseract':      ('tesseract（图片 OCR 引擎，可选）',   'apt install tesseract-ocr  /  dnf install tesseract'),
    'chi_sim':        ('tesseract 中文语言包 chi_sim', 'apt install tesseract-ocr-chi-sim  /  dnf install tesseract-langpack-chi_sim'),
    'pdftotext':      ('pdftotext（PDF 文本提取，可选）',     'apt install poppler-utils  /  dnf install poppler-utils'),
    'python-docx':    ('python-docx（读 Word）',        'pip install python-docx'),
    'openpyxl':       ('openpyxl（读 Excel）',          'pip install openpyxl'),
    'python-pptx':    ('python-pptx（读 PPT）',         'pip install python-pptx'),
    'pypdf':          ('pypdf（PDF 备用解析）',         'pip install pypdf'),
}

def _has_module(name):
    try:
        __import__(name)
        return True
    except ImportError:
        return False


def check_env(need_content=False):
    """返回缺失依赖的 key 列表。need_content=True 时检查全部（classify/auto 模式）

    关键：OCR 和 PDF 都是「两个方案任一可用即可」——
      OCR : rapidocr（纯 Python） 或 tesseract（系统程序）
      PDF : pypdf（纯 Python）    或 pdftotext（系统程序）
    所以只要 pip 装齐，就完全不依赖系统程序。
    """
    missing = []
    if not need_content:
        return missing

    # OCR
    has_rapid = _has_module('rapidocr_onnxruntime')
    has_tess = bool(shutil.which('tesseract'))
    if not has_rapid and not has_tess:
        missing.append('rapidocr')
    elif has_tess and not has_rapid:
        langs = _run(['tesseract', '--list-langs'], 15)
        if 'chi_sim' not in langs:
            missing.append('chi_sim')

    # PDF
    if not _has_module('pypdf') and not shutil.which('pdftotext'):
        missing.append('pypdf')

    for mod, key in [('docx', 'python-docx'), ('openpyxl', 'openpyxl'),
                     ('pptx', 'python-pptx')]:
        if not _has_module(mod):
            missing.append(key)
    return missing

def print_env_report(missing):
    if not missing:
        print('✅ 环境检查：内容识别所需依赖齐全\n')
        return
    print('⚠️  环境检查：以下依赖缺失 —— 对应格式将标为「未分类」（其它功能不受影响）')
    for k in missing:
        name, how = DEP_HINTS.get(k, (k, ''))
        print(f'   · {name}')
        if how:
            print(f'       装法：{how}')
    print('   提示：只做重命名不读内容的话，可以忽略以上缺失（用 rename 模式）\n')

# ============================================================================
# 三、文本提取（延迟导入 —— rename 模式不会加载这些重依赖）
# ============================================================================
def _run(cmd, timeout=60):
    try:
        r = subprocess.run(cmd, capture_output=True, timeout=timeout)
        return r.stdout.decode('utf-8', errors='ignore')
    except Exception:
        return ''

def extract_pdf(p):
    t = _run(['pdftotext', '-l', '10', str(p), '-'], 45)
    if t.strip():
        return t
    try:
        from pypdf import PdfReader
        r = PdfReader(str(p)); parts = []
        for i, pg in enumerate(r.pages):
            if i >= 10: break
            parts.append(pg.extract_text() or '')
        return "\n".join(parts)
    except Exception:
        return ''

def extract_docx(p):
    try:
        from docx import Document
        d = Document(str(p))
        parts = [x.text for x in d.paragraphs]
        for tb in d.tables:
            for row in tb.rows:
                parts.append(" ".join(c.text for c in row.cells))
        return "\n".join(parts)
    except Exception:
        return ''

def extract_xlsx(p):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets[:5]:
            for row in ws.iter_rows(max_row=300, values_only=True):
                parts.append(" ".join(str(c) for c in row if c is not None))
        wb.close()
        return "\n".join(parts)
    except Exception:
        return ''

def extract_pptx(p):
    try:
        from pptx import Presentation
        pr = Presentation(str(p)); parts = []
        for i, s in enumerate(pr.slides):
            if i >= 40: break
            for sh in s.shapes:
                if sh.has_text_frame:
                    parts.append(sh.text_frame.text)
        return "\n".join(parts)
    except Exception:
        return ''

def extract_plain(p):
    try:
        return Path(p).read_text(encoding='utf-8', errors='ignore')
    except Exception:
        return ''

_OCR_ENGINE = None


def _ocr_rapid(p):
    """rapidocr（纯 Python + onnxruntime，无需系统依赖，中文识别更准）"""
    global _OCR_ENGINE
    try:
        if _OCR_ENGINE is None:
            from rapidocr_onnxruntime import RapidOCR
            _OCR_ENGINE = RapidOCR()
        res, _ = _OCR_ENGINE(str(p))
        return '\n'.join(x[1] for x in res) if res else ''
    except Exception:
        return ''


def extract_image(p):
    # ① rapidocr 优先（pip 可装、能被 PyInstaller 打包、中文准确率明显更高）
    t = _ocr_rapid(p)
    if t.strip():
        return t
    # ② 回退系统 tesseract（装了的话）
    return _run(['tesseract', str(p), '-', '-l', 'chi_sim+eng'], 90)

EXTRACTORS = {
    '.pdf': extract_pdf, '.docx': extract_docx, '.xlsx': extract_xlsx, '.pptx': extract_pptx,
    '.txt': extract_plain, '.md': extract_plain,
    '.jpg': extract_image, '.jpeg': extract_image, '.png': extract_image,
}

def extract_text(p: Path) -> str:
    fn = EXTRACTORS.get(p.suffix.lower())
    if not fn: return ''
    try:
        return (fn(p) or '')[:MAX_CHARS]
    except Exception:
        return ''

# ============================================================================
# 四、内容分类（加权打分）
# ============================================================================
def classify(text: str):
    if not text.strip():
        return None, 0.0, []
    scores = {}
    for cat, items in _WEIGHTED.items():
        s = 0.0
        for kw, w in items:
            if kw in text:
                s += w + min(text.count(kw) - 1, 5) * 0.2
        scores[cat] = round(s, 1)
    ranked = sorted(scores.items(), key=lambda x: -x[1])
    top, ts = ranked[0]
    second = ranked[1][1] if len(ranked) > 1 else 0
    if ts >= 4.0 and ts > second:
        return top, ts, ranked[:3]
    return None, ts, ranked[:3]

# ============================================================================
# 五、文件名清理
# ============================================================================
JUNK_RE = re.compile(r'(最终版|最新版|修改版|修订版|定稿版|终版|的副本|副本|copy|final|new\s*\d*)\s*', re.I)

def clean_stem(stem, sep='_'):
    s = stem.strip()
    s = JUNK_RE.sub('', s)
    s = re.sub(r'[\s\u3000]+', sep, s)
    s = re.sub(r'[\\/:*?"<>|\u0000-\u001f]', '', s)
    s = re.sub(r'[()（）\[\]【】]', '', s)
    s = re.sub(r'(?:^|(?<=[_\-]))版(?=[_\-]|$)', '', s)
    s = re.sub(re.escape(sep) + r'{2,}', sep, s)
    s = re.sub(r'-{2,}', '-', s)
    return s.strip(sep + '-. ')

# ============================================================================
# 六、生成计划
# ============================================================================
def build_plan(directory, mode='name', sep='_', add_date=True, lower=False, pad=False):
    d = Path(directory)
    if not d.is_dir():
        print(f'错误：目录不存在 {directory}'); sys.exit(1)

    files = [p for p in sorted(d.iterdir())
             if p.is_file() and not p.name.startswith(('.', '_filekit_log_', '_rename_log_', '_classify_log_'))]
    if mode != 'name':
        print(f'正在读取 {len(files)} 个文件的内容...\n')

    plan = []
    for p in files:
        ext = p.suffix
        prefix = TYPE_PREFIX.get(ext.lower(), 'FILE')

        # ---- 分类（classify / auto 模式）----
        text, cat, score, top3 = '', None, 0.0, []
        if mode in ('classify', 'auto'):
            text = extract_text(p)
            if text.strip():
                cat, score, top3 = classify(text)
            elif mode == 'classify':
                pass                                  # 读不到内容 → 未分类
            # auto 模式：读不到内容 → 退化为纯文件名模式

        stem = p.stem
        has_prefix = bool(PREFIX_RE.match(stem))
        clean = clean_stem(stem, sep) or 'unnamed'

        # ---- 拼装 ----
        if mode == 'name' or (mode == 'auto' and not text.strip()):
            # 纯文件名模式：前缀 + 名字
            new_stem = clean if has_prefix else f'{prefix}{sep}{clean}'
        else:
            cat_part = cat if cat else '未分类'
            cat_in_name = cat_part in clean.split(sep)
            if not has_prefix:
                new_stem = f'{prefix}{sep}{clean}' if cat_in_name \
                           else f'{prefix}{sep}{cat_part}{sep}{clean}'
            else:
                new_stem = clean if cat_in_name else f'{clean}{sep}{cat_part}'

        if pad:
            new_stem = re.sub(r'(?<![A-Za-z])(\d+)$', lambda m: m.group(1).zfill(2), new_stem)
        if lower:
            new_stem = new_stem.lower()
        if add_date:
            mt = datetime.datetime.fromtimestamp(p.stat().st_mtime).strftime('%y%m%d%H%M')
            if mt not in new_stem:
                new_stem = f'{new_stem}{sep}{mt}'

        plan.append({'old': p.name, 'new': f'{new_stem}{ext}', 'cat': cat or '未分类',
                     'score': score, 'chars': len(text), 'top3': top3})

    # ---- 冲突处理 ----
    seen = {}
    for item in plan:
        base, ext = os.path.splitext(item['new'])
        name = item['new']; key = name.lower()
        while key in seen or ((d / name).exists() and name != item['old']):
            seen[key] = seen.get(key, 1) + 1
            name = f'{base}_{seen[key]:02d}{ext}'; key = name.lower()
        seen[key] = 1
        item['new'] = name

    return [i for i in plan if i['old'] != i['new']]

# ============================================================================
# 七、执行 / 回滚
# ============================================================================
def do_apply(directory, plan):
    d = Path(directory); done = []
    for i in plan:
        (d / i['old']).rename(d / i['new'])
        done.append({'old': i['old'], 'new': i['new'],
                     'category': i.get('cat'), 'score': i.get('score')})
    log_path = d / f'_filekit_log_{datetime.datetime.now():%Y%m%d_%H%M%S}.json'
    json.dump({'directory': str(d.resolve()), 'version': VERSION,
               'time': datetime.datetime.now().isoformat(timespec='seconds'), 'renames': done},
              open(log_path, 'w', encoding='utf-8'), ensure_ascii=False, indent=2)
    print(f'\n✅ 完成 {len(done)} 个')
    print(f'📄 日志：{log_path}')
    print(f'↩️  回滚：python3 {sys.argv[0]} rollback "{log_path}"')

def do_rollback(logfile):
    log = json.load(open(logfile, encoding='utf-8'))
    d = Path(log['directory']); n = 0
    for it in log['renames']:
        cur, old = d / it['new'], d / it['old']
        if cur.exists() and not old.exists():
            cur.rename(old); n += 1
    print(f'✅ 已回滚 {n} 个文件')

# ============================================================================
# 八、CLI
# ============================================================================
def add_common(ap):
    ap.add_argument('--dir', required=True, help='目标目录')
    ap.add_argument('--apply', action='store_true', help='真正执行（默认只预览）')
    ap.add_argument('--sep', default='_', help='分隔符，默认 _')
    ap.add_argument('--no-date', action='store_true', help='不加时间戳（默认加 YYMMDDHHMM）')
    ap.add_argument('--lower', action='store_true', help='文件名转小写')
    ap.add_argument('--pad', action='store_true', help='末尾数字补零')
    ap.add_argument('--verbose', action='store_true', help='显示分类候选得分')

def main():
    ap = argparse.ArgumentParser(
        prog='filekit',
        description=f'filekit v{VERSION} —— 文件整理工具箱（重命名 + 内容归类）',
        epilog='示例：python filekit.py auto --dir ./文件夹 --apply')
    sub = ap.add_subparsers(dest='cmd', required=True)

    p_r = sub.add_parser('rename', help='只按文件名规范（零依赖）')
    add_common(p_r)

    p_c = sub.add_parser('classify', help='读内容自动归类（38 类）')
    add_common(p_c)

    p_a = sub.add_parser('auto', help='智能：能读内容就归类，否则只按文件名')
    add_common(p_a)

    p_b = sub.add_parser('rollback', help='按日志回滚')
    p_b.add_argument('logfile', help='日志文件路径')

    sub.add_parser('doctor', help='检查环境依赖（装了哪些、缺哪些）')

    args = ap.parse_args()

    if args.cmd == 'doctor':
        print_env_report(check_env(need_content=True)); return

    if args.cmd == 'rollback':
        do_rollback(args.logfile); return

    mode = {'rename': 'name', 'classify': 'classify', 'auto': 'auto'}[args.cmd]
    if mode != 'name':
        print_env_report(check_env(need_content=True))
    plan = build_plan(args.dir, mode=mode, sep=args.sep,
                      add_date=not args.no_date, lower=args.lower, pad=args.pad)
    if not plan:
        print('没有需要改名的文件'); return

    # ---- 打印计划 ----
    if mode == 'name':
        print(f'{"原文件名":<40} →  新文件名')
        print('-' * 92)
        for i in plan:
            print(f'{i["old"][:38]:<40} →  {i["new"]}')
    else:
        print(f'{"原文件名":<34} {"识别为":<10} {"分":>6} {"字数":>6}   新文件名')
        print('-' * 118)
        for i in plan:
            print(f'{i["old"][:32]:<34} {i["cat"]:<10} {i["score"]:>6} {i["chars"]:>6}   {i["new"]}')
            if args.verbose and i['top3']:
                print(f'{"":<34} 候选: ' + ' | '.join(f'{c}:{s}' for c, s in i['top3']))
    print(f'\n共 {len(plan)} 个待重命名')

    # ---- 疑似同名检测 ----
    groups = {}
    for it in plan:
        base = re.sub(r'_\d{10}(_\d{2})?$', '', os.path.splitext(it['new'])[0])
        groups.setdefault(base, []).append(it['new'])
    dupes = {k: v for k, v in groups.items() if len(v) > 1}
    if dupes:
        print(f'\n⚠️  发现 {len(dupes)} 组疑似同名文档（可能是同一文档的多个版本）：')
        for k, v in dupes.items():
            print(f'  · {k}  →  {len(v)} 份')
            for name in v:
                print(f'        {name}')
        print('  → 时间戳已把版本区分开；如需精简，请确认后手动删除过期版本')

    if not args.apply:
        print('\n👀 预览模式。执行请加 --apply')
        return

    do_apply(args.dir, plan)

if __name__ == '__main__':
    main()
